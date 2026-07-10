from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


HEADERS = {
    "A": "week",
    "B": "date",
    "C": "day",
    "D": "workout",
    "E": "exercise",
    "F": "targetSets",
    "G": "targetReps",
    "H": "set1Weight",
    "I": "set1Reps",
    "J": "set1Rpe",
    "K": "set2Weight",
    "L": "set2Reps",
    "M": "set2Rpe",
    "N": "set3Weight",
    "O": "set3Reps",
    "P": "set3Rpe",
    "Q": "set4Weight",
    "R": "set4Reps",
    "S": "set4Rpe",
    "V": "done",
    "W": "notes",
}


def clean(value):
    if hasattr(value, "isoformat"):
        return value.isoformat()[:10]
    return value


def main() -> None:
    workbook_path = Path(sys.argv[1])
    output_path = Path(sys.argv[2])
    workbook = load_workbook(workbook_path, data_only=True)
    workout_log = workbook["Workout Log"]
    dashboard = workbook["Dashboard"]
    bodyweight = workbook["Bodyweight"]

    rows = []
    for row_number in range(6, workout_log.max_row + 1):
        exercise = workout_log[f"E{row_number}"].value
        if not exercise:
            continue

        record = {"id": f"wl-{row_number}"}
        for column, key in HEADERS.items():
            record[key] = clean(workout_log[f"{column}{row_number}"].value)
        rows.append(record)

    bodyweight_rows = []
    for row_number in range(7, bodyweight.max_row + 1):
        week = bodyweight[f"A{row_number}"].value
        target = bodyweight[f"B{row_number}"].value
        actual = bodyweight[f"C{row_number}"].value
        notes = bodyweight[f"D{row_number}"].value
        if week or target or actual or notes:
            bodyweight_rows.append(
                {
                    "week": clean(week),
                    "target": clean(target),
                    "actual": clean(actual),
                    "notes": clean(notes),
                }
            )

    payload = {
        "monthStart": clean(dashboard["B3"].value),
        "plan": rows,
        "bodyweight": bodyweight_rows,
    }

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        "export const seedData = "
        + json.dumps(payload, indent=2)
        + ";\n",
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()

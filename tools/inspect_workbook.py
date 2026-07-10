from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def cell_value(cell):
    value = cell.value
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def main() -> None:
    workbook_path = Path(sys.argv[1])
    workbook = load_workbook(workbook_path, data_only=False)
    calculated = load_workbook(workbook_path, data_only=True)

    summary = []
    for worksheet in workbook.worksheets:
        calculated_sheet = calculated[worksheet.title]
        formulas = []
        non_empty_rows = []

        for row in worksheet.iter_rows():
            values = [cell_value(cell) for cell in row]
            if any(value not in (None, "") for value in values):
                compact = []
                for cell in row:
                    if cell.value not in (None, ""):
                        compact.append(
                            {
                                "cell": cell.coordinate,
                                "value": cell_value(cell),
                                "calculated": cell_value(calculated_sheet[cell.coordinate]),
                            }
                        )
                non_empty_rows.append(compact)

            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    formulas.append({"cell": cell.coordinate, "formula": cell.value})

        summary.append(
            {
                "title": worksheet.title,
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
                "first_rows": non_empty_rows[:12],
                "formula_count": len(formulas),
                "formulas": formulas[:40],
            }
        )

    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
